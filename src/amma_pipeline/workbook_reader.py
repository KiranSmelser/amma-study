"""Read Amma Excel exports and demographic CSV files without mutating them."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .config import (
    DEMOGRAPHIC_REQUIRED_COLUMNS,
    REQUIRED_COLUMNS,
    REQUIRED_SHEETS,
)


class InputFormatError(ValueError):
    """Raised when a source file does not satisfy the required input contract."""


@dataclass(frozen=True)
class WorkbookData:
    """In-memory representation of the source workbook."""

    tables: dict[str, list[dict[str, Any]]]
    extract_metadata: dict[str, Any]


def _clean_header(value: object) -> str:
    return str(value or "").lstrip("\ufeff").strip()


def _table_records(worksheet: Any) -> list[dict[str, Any]]:
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [_clean_header(value) for value in rows[0]]
    if not any(headers):
        return []
    if any(not header for header in headers):
        raise InputFormatError(f"Blank header in sheet {worksheet.title!r}")
    if len(headers) != len(set(headers)):
        raise InputFormatError(f"Duplicate header in sheet {worksheet.title!r}")

    records: list[dict[str, Any]] = []
    for row_number, row in enumerate(rows[1:], start=2):
        if not any(value not in (None, "") for value in row):
            continue
        padded = tuple(row) + (None,) * max(0, len(headers) - len(row))
        record = dict(zip(headers, padded, strict=False))
        record["__source_row__"] = row_number
        records.append(record)
    return records


def read_workbook(path: str | Path) -> WorkbookData:
    """Load and validate the structural contract of an Amma workbook."""

    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(source)

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        missing_sheets = [name for name in REQUIRED_SHEETS if name not in workbook]
        if missing_sheets:
            raise InputFormatError(
                "Missing required workbook sheet(s): " + ", ".join(missing_sheets)
            )

        tables: dict[str, list[dict[str, Any]]] = {}
        for sheet_name, required_headers in REQUIRED_COLUMNS.items():
            records = _table_records(workbook[sheet_name])
            actual_headers = (
                set(records[0]) - {"__source_row__"}
                if records
                else {
                    _clean_header(cell.value)
                    for cell in next(
                        workbook[sheet_name].iter_rows(min_row=1, max_row=1)
                    )
                }
            )
            missing_headers = [
                header for header in required_headers if header not in actual_headers
            ]
            if missing_headers:
                raise InputFormatError(
                    f"Sheet {sheet_name!r} is missing column(s): "
                    + ", ".join(missing_headers)
                )
            tables[sheet_name] = records

        extract_metadata: dict[str, Any] = {}
        for row in workbook["Extract info"].iter_rows(values_only=True):
            key = _clean_header(row[0] if row else None)
            if key:
                extract_metadata[key] = row[1] if len(row) > 1 else None

        return WorkbookData(tables=tables, extract_metadata=extract_metadata)
    finally:
        workbook.close()


def read_demographics(path: str | Path) -> list[dict[str, Any]]:
    """Read the restricted demographic input while handling a UTF-8 BOM."""

    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(source)

    with source.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        headers = [_clean_header(value) for value in (reader.fieldnames or [])]
        missing_headers = [
            header for header in DEMOGRAPHIC_REQUIRED_COLUMNS if header not in headers
        ]
        if missing_headers:
            raise InputFormatError(
                "Demographic CSV is missing column(s): " + ", ".join(missing_headers)
            )

        records: list[dict[str, Any]] = []
        for row_number, source_row in enumerate(reader, start=2):
            cleaned = {
                _clean_header(key): value for key, value in source_row.items() if key
            }
            if not any(value not in (None, "") for value in cleaned.values()):
                continue
            cleaned["__source_row__"] = row_number
            records.append(cleaned)
        return records
